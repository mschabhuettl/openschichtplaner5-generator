"""Portable synthetic regressions for DATA-01 through DATA-04."""
from collections import defaultdict
import csv
from datetime import date, datetime, timedelta

import pytest
from openpyxl import load_workbook

from sp5generator.domain import input_diagnostics, snapshot_hash
from sp5generator.export import export_table
from sp5generator.history_approvals import apply_history_approvals
from sp5generator.models import Assignment, BoundaryWork, Interval, Result
from sp5generator.project_creation import ProjectCreateRequest, create_project
from sp5generator.replacement import replacement_candidates
from sp5generator.solver import solve
from sp5generator.sp5_adapter import historical_matrix, import_snapshot
from sp5generator.validator import validate


@pytest.fixture
def native_source(tmp_path):
    """Real native facade and calculations over synthetic in-memory tables."""
    database = pytest.importorskip("sp5lib.database")

    class ArtificialNativeDatabase(database.SP5Database):
        def __init__(self, tables):
            super().__init__(str(tmp_path / "nonexistent-db"))
            self.tables = tables

        def _read(self, table):
            return [dict(row) for row in self.tables.get(table, [])]

        def _read_by_month(self, name, date_field="DATE"):
            grouped = defaultdict(list)
            for row in self._read(name):
                grouped[row[date_field][:7]].append(row)
            return grouped

    services = [{"ID": sid, "NAME": f"Synthetic service {sid}",
                 **{f"STARTEND{i}": "08:00-16:00" for i in range(8)},
                 **{f"DURATION{i}": 8.0 for i in range(8)}} for sid in (201, 202)]
    days = [date(2026, 1, 5) + timedelta(days=i) for i in range(3)]
    tables = {
        "EMPL": [{"ID": 101, "NAME": "Synthetic native person", "EMPSTART": "2025-01-01",
                  "EMPEND": "2027-12-31", "HRSDAY": 8.0, "HRSWEEK": 40.0,
                  "WORKDAYS": "1111100", "CALCBASE": 0}],
        "GROUP": [{"ID": 1, "NAME": "Synthetic group", "SUPERID": 0}],
        "GRASG": [{"ID": 1, "GROUPID": 1, "EMPLOYEEID": 101}],
        "SHIFT": services,
        "WOPL": [{"ID": 301, "NAME": "Synthetic workplace"}],
        "MASHI": [{"ID": i + 1, "EMPLOYEEID": 101, "DATE": str(day), "SHIFTID": 201,
                   "WORKPLACID": 301, "TYPE": 0} for i, day in enumerate(days)],
        "SPSHI": [{"ID": i + 1, "EMPLOYEEID": 101, "DATE": str(day), "SHIFTID": 202,
                   "WORKPLACID": 301, "TYPE": 0, "STARTEND": "08:00-16:00", "DURATION": 8.0}
                  for i, day in enumerate(days)],
    }
    return ArtificialNativeDatabase(tables), days


def test_data01_history_uses_effective_ist_for_approvals_and_demand(native_source):
    from sp5lib import calculations as calc

    db, days = native_source
    snapshot = import_snapshot(db, date(2026, 2, 2), date(2026, 2, 4), team_id="1", timezone="UTC",
                               demand_source="history", history_start=days[0], history_end=days[-1])
    matrix = historical_matrix(db, snapshot, days[0], days[-1], history_plan="ist")
    snapshot.metadata["history_matrix"] = matrix
    assert not snapshot.employees[0].approvals  # Evidence alone grants nothing.
    apply_history_approvals(snapshot, minimum_days=3)
    source_hours = calc.get_work_hours(
        calc.EmployeeContext.from_record(db.tables["EMPL"][0]), days[0], days[-1], holidays={},
        shifts_by_id={s["ID"]: s for s in db.tables["SHIFT"]},
        manual_shifts=db.tables["MASHI"], special_shifts=db.tables["SPSHI"],
    )
    assert source_hours == 24.0
    assert [a.function_id for a in snapshot.employees[0].approvals] == ["sp5:service:202"]
    assert matrix[0]["observed_assignment_count"] == 3
    assert matrix[0]["suggested_approvals"][0]["evidence_days"] == 3
    assert snapshot.metadata["history_demand"]["slots"] == 3
    assert sum(s.paid_minutes for s in snapshot.shifts) == 1440


@pytest.mark.parametrize("special", [False, True])
def test_data01_both_keeps_soll_and_accounts_for_each_source(native_source, special):
    db, days = native_source
    db.tables["MASHI"] += [{**row, "ID": row["ID"] + 10, "TYPE": 1}
                           for row in list(db.tables["MASHI"])]
    if not special:
        db.tables["SPSHI"] = []
    snapshot = import_snapshot(db, date(2026, 2, 2), date(2026, 2, 4), team_id="1", timezone="UTC")
    matrix = historical_matrix(db, snapshot, days[0], days[-1], history_plan="both")
    person, = matrix
    assert person["observed_assignment_count"] == 6
    observations = {row["shift_id"]: row for row in person["observed_shifts"]}
    assert observations["201"]["source_counts"] == ({"soll": 3} if special else {"ist": 3, "soll": 3})
    if special:
        assert observations["202"]["source_counts"] == {"special_shift": 3}
    assert all(a["evidence_days"] == 3 for a in person["suggested_approvals"])


def test_data01_soll_is_not_suppressed_by_special_work(native_source):
    db, days = native_source
    for row in db.tables["MASHI"]:
        row["TYPE"] = 1
    snapshot = import_snapshot(db, date(2026, 2, 2), date(2026, 2, 4), team_id="1", timezone="UTC",
                               reference_plan="soll", demand_source="history",
                               history_start=days[0], history_end=days[-1])
    matrix = historical_matrix(db, snapshot, days[0], days[-1], history_plan="soll")
    assert {s["shift_id"] for s in matrix[0]["observed_shifts"]} == {"201", "202"}
    assert matrix[0]["observed_assignment_count"] == 6
    assert snapshot.metadata["history_demand"]["slots"] == 6
    assert not snapshot.employees[0].approvals


@pytest.mark.parametrize("shift_id,special_type", [(202, 0), (202, 1), (0, 0)])
def test_data01_replacement_is_person_day_wide_not_workplace_or_type_specific(native_source, shift_id, special_type):
    db, days = native_source
    db.tables["EMPL"].append({**db.tables["EMPL"][0], "ID": 102, "NAME": "Other synthetic person"})
    db.tables["GRASG"].append({"ID": 2, "GROUPID": 1, "EMPLOYEEID": 102})
    db.tables["MASHI"] += [{**row, "ID": row["ID"] + 10, "EMPLOYEEID": 102}
                           for row in list(db.tables["MASHI"])]
    for row in db.tables["SPSHI"]:
        row.update(SHIFTID=shift_id, TYPE=special_type, WORKPLACID=0)
    snapshot = import_snapshot(db, date(2026, 2, 2), date(2026, 2, 4), team_id="1", timezone="UTC",
                               demand_source="history", history_start=days[0], history_end=days[-1])
    matrix = {row["employee_id"]: row for row in historical_matrix(db, snapshot, days[0], days[-1])}
    observed = {row["shift_id"] for row in matrix["sp5:employee:101"]["observed_shifts"]}
    assert observed == ({"201"} if not shift_id else {"202"} if special_type == 0 else set())
    assert matrix["sp5:employee:102"]["observed_assignment_count"] == 3
    assert matrix["sp5:employee:102"]["observed_shifts"][0]["shift_id"] == "201"
    assert snapshot.metadata["history_demand"]["slots"] == (3 if special_type == 1 else 6)


def project(*, people=2, start="2026-01-05", end="2026-01-06"):
    return create_project(ProjectCreateRequest(
        project_name="Synthetic data contracts", period_start=start, period_end=end, timezone="UTC",
        people=[{"name": f"Artificial person {i + 1}", "target_hours": 16.0} for i in range(people)],
        positions=[{"name": "Artificial function"}],
        shift_templates=[{"name": "Day shift", "kind": "day", "start_time": "08:00", "end_time": "16:00",
                          "weekdays": list(range(7)), "demands": [{"position": 0, "minimum": 1, "maximum": 1}]}],
        rules={"min_rest_hours": 11.0, "max_consecutive_work_days": 6, "max_consecutive_nights": 3,
               "max_daily_hours": 12.0, "max_weekly_hours": 48.0, "weekly_rest_hours": 0.0},
        rules_confirmed=True, approvals_confirmed=True, context_duty_free_confirmed=True,
    ))


@pytest.mark.parametrize("untimed", [False, True], ids=["timed", "untimed"])
def test_data02_exports_include_personal_work_and_balances(tmp_path, untimed):
    snapshot = project(people=1)
    employee = snapshot.employees[0]
    first, second = snapshot.shifts
    work = BoundaryWork(
        id="synthetic-personal-work", employee_id=employee.id,
        segments=[] if untimed else first.segments, day=snapshot.period_start if untimed else None,
        kind="unknown" if untimed else "day", in_period=True, paid_minutes=480,
    )
    snapshot.boundary_work = [work]
    snapshot.shifts = [second]
    snapshot.demands = [snapshot.demands[1]]
    assert input_diagnostics(snapshot) == []
    result = solve(snapshot, time_limit=5, workers=1)
    assert result.validation.valid and result.validation.complete
    assert result.metrics["employees"][employee.id]["paid_minutes"] == 960
    csv_path, xlsx_path = tmp_path / "personal.csv", tmp_path / "personal.xlsx"
    export_table(snapshot, result, csv_path)
    export_table(snapshot, result, xlsx_path)
    with csv_path.open(encoding="utf-8-sig", newline="") as stream:
        csv_rows = list(csv.reader(stream))
    totals = next(row for row in csv_rows if len(row) == 4 and row[0] == employee.id)
    assert (int(totals[2]), int(totals[3])) == (960, 0)
    personal_rows = [row for row in csv_rows if len(row) >= 10 and row[2] == work.id]
    assert len(personal_rows) == 1
    assert personal_rows[0][8:10] == ["480", "True"]
    book = load_workbook(xlsx_path, data_only=True)
    try:
        balances = book["Stundenübersicht"]
        assert (balances["C5"].value, balances["F5"].value, balances["G5"].value) == (16, 0, 2)
        calendar = book["Plan 2026-01"]["B5"].value
        assert calendar and "Persönliche Arbeit" in calendar
        details = [row for row in book["Einteilungen"].values if len(row) >= 10 and row[2] == work.id]
        assert len(details) == 1 and details[0][8:10] == (480, True)
        if untimed:
            assert personal_rows[0][6:8] == ["", ""]
            assert details[0][6:8] == (None, None)
            assert str(work.day) in personal_rows[0]
            assert "ohne Zeiten" in calendar and "00:00" not in calendar
        else:
            assert personal_rows[0][6:8] == [first.segments[0].start.isoformat(), first.segments[0].end.isoformat()]
    finally:
        book.close()


def test_data02_split_overnight_work_counts_once_and_context_has_no_period_credit(tmp_path):
    snapshot = project(people=1, end="2026-01-07")
    employee = snapshot.employees[0]
    employee.credit_minutes, employee.balance_minutes = 60, -120

    def interval(start, end):
        return Interval(start=datetime.fromisoformat(start), end=datetime.fromisoformat(end))

    snapshot.boundary_work = [
        BoundaryWork(id="context", employee_id=employee.id, kind="night", paid_minutes=480,
                     segments=[interval("2026-01-04T22:00+00:00", "2026-01-05T06:00+00:00")]),
        BoundaryWork(id="split-personal", employee_id=employee.id, kind="night", paid_minutes=420,
                     in_period=True, segments=[
                         interval("2026-01-05T22:00+00:00", "2026-01-06T02:00+00:00"),
                         interval("2026-01-06T03:00+00:00", "2026-01-06T06:00+00:00"),
                     ]),
    ]
    snapshot.metadata["provenance"] = {"split-personal": {
        "name": "=Synthetic personal service", "function_id": "synthetic-function", "workplace_id": 301,
    }}
    snapshot.demands, snapshot.shifts = snapshot.demands[2:], snapshot.shifts[2:]
    result = solve(snapshot, time_limit=5, workers=1)
    assert result.validation.complete and result.validation.valid
    assert result.metrics["employees"][employee.id]["paid_minutes"] == 900
    export_table(snapshot, result, tmp_path / "split.csv")
    export_table(snapshot, result, tmp_path / "split.xlsx")
    with (tmp_path / "split.csv").open(encoding="utf-8-sig", newline="") as stream:
        csv_rows = list(csv.reader(stream))
    totals = next(row for row in csv_rows if len(row) == 4 and row[0] == employee.id)
    assert tuple(map(int, totals[1:])) == (960, 960, -120)
    detail = [row for row in csv_rows if len(row) >= 10 and row[2] == "split-personal"]
    assert len(detail) == 2 and [row[8] for row in detail] == ["420", ""]
    assert all(row[3:6] == ["synthetic-function", "301", "'=Synthetic personal service"] for row in detail)
    book = load_workbook(tmp_path / "split.xlsx", data_only=True)
    try:
        hours = book["Stundenübersicht"]
        assert (hours["C5"].value, hours["F5"].value, hours["G5"].value, hours["H5"].value) == (15, -2, 3, 1)
        calendar = book["Plan 2026-01"]
        assert "Randdienst" in calendar["B5"].value
        assert "22:00–02:00 (+1)" in calendar["B5"].value
        assert "22:00–02:00 (+1)" in calendar["C5"].value
        assert "03:00–06:00" in calendar["C5"].value
        assert all(cell.data_type != "f" for sheet in book for row in sheet for cell in row)
        xlsx_detail = [row for row in book["Einteilungen"].values if row[2] == "split-personal"]
        assert len(xlsx_detail) == 2 and [row[8] for row in xlsx_detail] == [420, None]
    finally:
        book.close()


@pytest.mark.parametrize("suffix", ["csv", "xlsx"])
def test_data02_duplicate_personal_and_staffing_work_is_rejected(tmp_path, suffix):
    snapshot = project(people=1)
    result = solve(snapshot, time_limit=5, workers=1)
    snapshot.boundary_work = [BoundaryWork(
        id="duplicate-personal", employee_id=snapshot.employees[0].id,
        kind="day", segments=snapshot.shifts[0].segments, paid_minutes=480, in_period=True,
    )]
    result.snapshot_hash = snapshot_hash(snapshot)
    assert not validate(snapshot, result.assignments).valid
    target = tmp_path / ("duplicate." + suffix)
    with pytest.raises(ValueError, match="Ungültigen Entwurf"):
        export_table(snapshot, result, target)
    assert not target.exists()


@pytest.mark.parametrize("begin,end,expected", [
    ("2026-01-04T20:00+00:00", "2026-01-05T00:30+00:00", 0),
    ("2026-01-04T23:30+00:00", "2026-01-05T01:30+00:00", 120),
])
def test_data02_period_credit_uses_local_start_day_not_marker_or_utc_date(tmp_path, begin, end, expected):
    """Check exports of a real independently validated, empty staffing proposal."""
    snapshot = project(people=1, end="2026-01-05")
    snapshot.timezone = "Europe/Vienna"
    snapshot.demands, snapshot.shifts = [], []
    employee = snapshot.employees[0]
    snapshot.boundary_work = [BoundaryWork(
        id="local-start", employee_id=employee.id, kind="day", paid_minutes=120,
        in_period=True, segments=[Interval(start=datetime.fromisoformat(begin), end=datetime.fromisoformat(end))],
    )]
    checked = validate(snapshot, [])
    assert checked.valid and checked.complete
    result = Result(snapshot_id=snapshot.id, snapshot_hash=snapshot_hash(snapshot),
                    solver_status="FEASIBLE", validation=checked, runtime_seconds=0, assignments=[])
    export_table(snapshot, result, tmp_path / "local.csv")
    export_table(snapshot, result, tmp_path / "local.xlsx")
    with (tmp_path / "local.csv").open(encoding="utf-8-sig", newline="") as stream:
        totals = next(row for row in csv.reader(stream) if len(row) == 4 and row[0] == employee.id)
    assert int(totals[2]) == expected
    book = load_workbook(tmp_path / "local.xlsx", data_only=True)
    try:
        assert book["Stundenübersicht"]["C5"].value == expected / 60
        assert book["Stundenübersicht"]["G5"].value == 1
        assert book["Plan 2026-01"]["B5"].value
    finally:
        book.close()


def replace_duties(original, absent, candidate, demand_ids):
    return [a.model_copy(update={"employee_id": candidate})
            if a.employee_id == absent and a.demand_id in demand_ids else a
            for a in original]


@pytest.mark.parametrize("limit,value,code", [
    ("max_weekly_minutes", 480, "weekly_limit"),
    ("max_period_minutes", 480, "period_limit"),
    ("max_work_days", 1, "work_days"),
    ("max_consecutive_work_days", 1, "consecutive_work"),
])
def test_data03_single_candidate_respects_cumulative_weekly_limit(limit, value, code):
    snapshot = project()
    absent, candidate = snapshot.employees
    setattr(snapshot.profiles[0], limit, value)
    first, second = snapshot.demands
    original = [Assignment(employee_id=candidate.id, demand_id=first.id),
                Assignment(employee_id=absent.id, demand_id=second.id)]
    assert validate(snapshot, original).complete
    hypothetical = replace_duties(original, absent.id, candidate.id, {second.id})
    invalid = validate(snapshot, hypothetical)
    assert not invalid.valid and code in {d.code for d in invalid.diagnostics}
    before = snapshot.model_dump(), [a.model_dump() for a in original]
    report = replacement_candidates(snapshot, original, absent.id, snapshot.period_end, snapshot.period_end)
    assert candidate.id not in [c["employee_id"] for c in report["duties"][0]["candidates"]]
    assert report["duties"][0]["blocked"] == {code: 1}
    assert candidate.id not in report["covers_whole_absence"]
    assert (snapshot.model_dump(), [a.model_dump() for a in original]) == before


@pytest.mark.parametrize("invalid_input", [False, True])
def test_data03_candidate_uses_validator_validity_not_just_diagnostic_labels(invalid_input):
    snapshot = project()
    absent, candidate = snapshot.employees
    original = [Assignment(employee_id=absent.id, demand_id=d.id) for d in snapshot.demands]
    if invalid_input:
        snapshot.boundary_work = [BoundaryWork(
            id="outside-context", employee_id=absent.id, kind="day",
            segments=[i.model_copy(update={"start": i.start - timedelta(days=100),
                                           "end": i.end - timedelta(days=100)})
                      for i in snapshot.shifts[0].segments],
        )]
    else:
        snapshot.context_complete = False
    baseline = validate(snapshot, original)
    assert {d.code for d in baseline.diagnostics} == {"context"}
    assert baseline.valid is not invalid_input
    report = replacement_candidates(snapshot, original, absent.id, snapshot.period_start, snapshot.period_end)
    for duty in report["duties"]:
        assert [c["employee_id"] for c in duty["candidates"]] == ([] if invalid_input else [candidate.id])
    assert report["covers_whole_absence"] == ([] if invalid_input else [candidate.id])


def test_data03_whole_absence_requires_joint_validation():
    snapshot = project()
    absent, candidate = snapshot.employees
    limited = snapshot.profiles[0].model_copy(deep=True, update={"id": "candidate-rules", "max_weekly_minutes": 480})
    snapshot.profiles.append(limited)
    candidate.profile_ids = [limited.id]
    original = [Assignment(employee_id=absent.id, demand_id=d.id) for d in snapshot.demands]
    assert validate(snapshot, original).complete
    for demand in snapshot.demands:
        assert validate(snapshot, replace_duties(original, absent.id, candidate.id, {demand.id})).complete
    combined = replace_duties(original, absent.id, candidate.id, {d.id for d in snapshot.demands})
    invalid = validate(snapshot, combined)
    assert not invalid.valid and "weekly_limit" in {d.code for d in invalid.diagnostics}
    report = replacement_candidates(snapshot, original, absent.id, snapshot.period_start, snapshot.period_end)
    assert all(candidate.id in {c["employee_id"] for c in duty["candidates"]} for duty in report["duties"])
    assert candidate.id not in report["covers_whole_absence"]


def test_data03_vacancies_do_not_block_valid_individual_or_joint_candidates():
    snapshot = project(end="2026-01-07")
    absent, candidate = snapshot.employees
    original = [Assignment(employee_id=absent.id, demand_id=d.id) for d in snapshot.demands[:2]]
    baseline = validate(snapshot, original)
    assert baseline.valid and not baseline.complete
    assert {d.code for d in baseline.diagnostics} == {"vacancy"}
    report = replacement_candidates(snapshot, original, absent.id, snapshot.period_start, snapshot.period_end)
    assert len(report["duties"]) == 2
    assert all([c["employee_id"] for c in duty["candidates"]] == [candidate.id] for duty in report["duties"])
    assert report["covers_whole_absence"] == [candidate.id]
    hypothetical = replace_duties(original, absent.id, candidate.id, {d.id for d in snapshot.demands[:2]})
    assert validate(snapshot, hypothetical).valid


@pytest.mark.parametrize("same_day", [False, True], ids=["previous-day", "same-day"])
def test_data04_replacement_preserves_untimed_personal_work(same_day):
    snapshot = project()
    absent, candidate = snapshot.employees
    work = BoundaryWork(
        id="untimed-personal", employee_id=candidate.id, segments=[],
        day=snapshot.period_end if same_day else snapshot.period_start,
        in_period=True, paid_minutes=480,
    )
    snapshot.boundary_work = [work]
    snapshot.demands, snapshot.shifts = snapshot.demands[1:], snapshot.shifts[1:]
    original = [Assignment(employee_id=absent.id, demand_id=snapshot.demands[0].id)]
    assert validate(snapshot, original).complete
    before = snapshot.model_dump(), [a.model_dump() for a in original]
    report = replacement_candidates(snapshot, original, absent.id, snapshot.period_end, snapshot.period_end)
    duty, = report["duties"]
    if same_day:
        assert duty["candidates"] == []
        assert duty["blocked"] == {"personal_work": 1}
        assert report["covers_whole_absence"] == []
    else:
        assert duty["candidates"] == [{"employee_id": candidate.id, "days_since_last_duty": 1}]
        assert report["covers_whole_absence"] == [candidate.id]
        assert validate(snapshot, replace_duties(original, absent.id, candidate.id, {snapshot.demands[0].id})).complete
    assert (snapshot.model_dump(), [a.model_dump() for a in original]) == before
