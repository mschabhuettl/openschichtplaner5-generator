"""Portable synthetic regressions from the independent DATA review."""
import csv
from datetime import UTC, date, datetime, timedelta

import pytest
from openpyxl import load_workbook

from sp5generator.domain import input_diagnostics, snapshot_hash
from sp5generator.export import export_table
from sp5generator.models import (
    Approval, Assignment, BoundaryWork, Demand, Employee, Interval, Position, Result,
    RuleProfile, Shift, Snapshot,
)
from sp5generator.replacement import replacement_candidates
from sp5generator.solver import solve
from sp5generator.validator import validate


def snapshot_for_review(*, night=False, days=3):
    first = date(2026, 1, 5)
    last = first + timedelta(days=days - 1)
    context_start, context_end = first - timedelta(days=10), last + timedelta(days=10)
    profile = RuleProfile(
        id="review-rules", valid_from=context_start, valid_until=context_end,
        min_rest_minutes=660, confirmed=True, source="synthetic-review",
    )
    position = Position(
        id="review-position", name="Synthetic role", function_id="review-function",
        workplace_id="review-workplace", qualifications_required=False,
    )
    people = [Employee(
        id=pid, name=pid, team_ids=["review-team"],
        employment_start=context_start, employment_end=context_end,
        profile_ids=[profile.id], target_minutes=1000,
        approvals=[Approval(function_id=position.function_id,
                            workplace_id=position.workplace_id,
                            valid_from=context_start, valid_until=context_end)],
    ) for pid in ("absent", "candidate")]
    shifts = []
    for index in range(days):
        start = datetime.combine(first + timedelta(days=index), datetime.min.time(), UTC)
        start += timedelta(hours=20 if night else 8)
        shifts.append(Shift(
            id=f"review-shift-{index}", name="Synthetic duty", kind="night" if night else "day",
            team_id="review-team", segments=[Interval(start=start, end=start + timedelta(hours=10 if night else 8))],
            paid_minutes=600 if night else 480,
        ))
    return Snapshot(
        id="review-snapshot", revision="1", created_at=datetime(2026, 1, 1, tzinfo=UTC),
        timezone="UTC", period_start=first, period_end=last,
        context_start=context_start, context_end=context_end, context_complete=True,
        rule_version="1", source="synthetic", employees=people, positions=[position],
        shifts=shifts, demands=[Demand(id=f"review-demand-{i}", shift_id=s.id,
                                     position_id=position.id, minimum=1, maximum=1)
                                 for i, s in enumerate(shifts)], profiles=[profile],
    )


def test_joint_night_block_bridge_is_recommended():
    snapshot = snapshot_for_review(night=True)
    snapshot.profiles[0].after_night_block_rest_minutes = 2880
    snapshot.profiles[0].night_block_gap_days = 1
    original = [Assignment(employee_id="absent" if i < 2 else "candidate", demand_id=d.id)
                for i, d in enumerate(snapshot.demands)]
    assert validate(snapshot, original).complete
    first_only = [a.model_copy(update={"employee_id": "candidate"})
                  if a.demand_id == snapshot.demands[0].id else a for a in original]
    single = validate(snapshot, first_only)
    assert not single.valid
    assert {d.code for d in single.diagnostics} == {"night_block"}
    # Jan 6 bridges the Jan 5/7 blocks; joint feasibility is not monotonic.
    all_transferred = [a.model_copy(update={"employee_id": "candidate"}) for a in original]
    combined = validate(snapshot, all_transferred)
    assert combined.valid and combined.complete
    before = snapshot.model_dump(), [a.model_dump() for a in original]
    report = replacement_candidates(snapshot, original, "absent", snapshot.period_start,
                                    snapshot.period_start + timedelta(days=1))
    assert before == (snapshot.model_dump(), [a.model_dump() for a in original])
    assert report["duties"][0]["candidates"] == []
    assert report["covers_whole_absence"] == ["candidate"]


def checked_result(snapshot, assignments):
    checked = validate(snapshot, assignments)
    assert checked.valid, checked.model_dump()
    return Result(snapshot_id=snapshot.id, snapshot_hash=snapshot_hash(snapshot),
                  solver_status="FEASIBLE", assignments=assignments,
                  validation=checked, runtime_seconds=0)


@pytest.mark.parametrize("invalid", [None, [], "not a mapping", 17, False])
@pytest.mark.parametrize("level", ["provenance", "entry"])
@pytest.mark.parametrize("suffix", ["csv", "xlsx"])
def test_optional_nonmapping_provenance_exports_with_literal_ids(tmp_path, invalid, level, suffix):
    snapshot = snapshot_for_review(days=1)
    snapshot.shifts, snapshot.demands = [], []
    work_id = " personal\u00a0ID "
    snapshot.boundary_work = [BoundaryWork(
        id=work_id, employee_id="candidate", segments=[], day=snapshot.period_start,
        in_period=True, paid_minutes=480,
    )]
    snapshot.metadata = {"provenance": invalid if level == "provenance" else {work_id: invalid}}
    snapshot = Snapshot.model_validate(snapshot.model_dump())
    assert input_diagnostics(snapshot) == []
    result = checked_result(snapshot, [])
    before = snapshot.model_dump()
    target = tmp_path / ("synthetic." + suffix)
    target.write_bytes(b"previous synthetic export")
    export_table(snapshot, result, target)
    if suffix == "csv":
        with target.open(encoding="utf-8-sig", newline="") as stream:
            detail = [row for row in csv.reader(stream) if len(row) == 12 and row[2] == work_id]
    else:
        book = load_workbook(target)
        try:
            detail = [["" if value is None else str(value) for value in row]
                      for row in book["Einteilungen"].values if row[2] == work_id]
        finally:
            book.close()
    assert len(detail) == 1
    assert detail[0][2:6] == [work_id, "", "", "Persönliche Arbeit"]
    assert detail[0][8] == "480"
    assert snapshot.model_dump() == before


@pytest.mark.parametrize("exists", [False, True], ids=["new", "existing"])
@pytest.mark.parametrize("failure", ["row-generation", "encoding"])
def test_csv_generation_failure_preserves_destination(tmp_path, monkeypatch, exists, failure):
    import sp5generator.export as export_module

    snapshot = snapshot_for_review(days=1)
    assignments = [Assignment(employee_id="candidate", demand_id=snapshot.demands[0].id)]
    if failure == "encoding":
        # Accepted source text which UTF-8 cannot encode; fail after headers.
        snapshot.employees[1].name = "Synthetic invalid Unicode \ud800"
        error = UnicodeEncodeError
    else:
        original_rows = export_module.rows

        def fail_after_header(snapshot, result):
            yield next(original_rows(snapshot, result))
            raise RuntimeError("synthetic row generation failure")

        monkeypatch.setattr(export_module, "rows", fail_after_header)
        error = RuntimeError
    assert input_diagnostics(snapshot) == []
    result = checked_result(snapshot, assignments)
    target = tmp_path / "synthetic.csv"
    previous = b"previous synthetic export\r\n"
    if exists:
        target.write_bytes(previous)
    with pytest.raises(error):
        export_table(snapshot, result, target)
    if exists:
        assert target.read_bytes() == previous
    else:
        assert not target.exists()
    assert set(tmp_path.iterdir()) == ({target} if exists else set())


def test_csv_destination_is_published_only_after_generation(tmp_path, monkeypatch):
    import sp5generator.export as export_module

    snapshot = snapshot_for_review(days=1)
    result = checked_result(snapshot, [Assignment(employee_id="candidate", demand_id=snapshot.demands[0].id)])
    target = tmp_path / "synthetic.csv"
    previous = b"previous synthetic export"
    target.write_bytes(previous)
    original_rows = export_module.rows
    observed = []

    def observe_generation(snapshot, result):
        for row in original_rows(snapshot, result):
            observed.append(target.read_bytes())
            yield row

    monkeypatch.setattr(export_module, "rows", observe_generation)
    export_table(snapshot, result, target)
    assert observed and all(data == previous for data in observed)
    assert target.read_bytes().startswith(b"\xef\xbb\xbfZeitraum,")
    with target.open(encoding="utf-8-sig", newline="") as stream:
        totals = next(row for row in csv.reader(stream) if len(row) == 4 and row[0] == "candidate")
    assert totals[2] == "480"
    assert set(tmp_path.iterdir()) == {target}


PERIOD_CASES = [
    ("prior-marked-in-period", 0),
    ("untimed-default-marker", 120),
    ("timed-marked-in-period", 120),
    ("local-midnight-marked-in-period", 120),
    ("prior-tail-marked-in-period", 0),
    ("future-marked-in-period", 0),
    ("split-reversed-marked-in-period", 120),
]


def accounting_snapshot(mode, *, staffing=False):
    snapshot = snapshot_for_review(days=2 if staffing else 1)
    snapshot.employees = [snapshot.employees[1]]
    snapshot.profiles[0].min_rest_minutes = 0  # Isolate paid-time accounting.
    snapshot.shifts = snapshot.shifts[1:] if staffing else []
    snapshot.demands = snapshot.demands[1:] if staffing else []
    kwargs: dict = {"segments": []}
    if mode == "untimed-default-marker":
        kwargs["day"] = snapshot.period_start
    else:
        # Timed in-period work requires explicit confirmation in the accepted
        # input contract. Untimed day-bearing work does not require the marker.
        kwargs.update(kind="day", in_period=True)
        start = datetime(2026, 1, 5, 8, tzinfo=UTC)
        end = start + timedelta(hours=2)
        if mode in {"prior-marked-in-period", "prior-tail-marked-in-period"}:
            start = datetime(2026, 1, 4, 20, tzinfo=UTC)
            end = start + timedelta(hours=6 if "tail" in mode else 2)
            kwargs["in_period"] = True
        elif mode == "future-marked-in-period":
            start = datetime.combine(snapshot.period_end + timedelta(days=1), datetime.min.time(), UTC)
            end = start + timedelta(hours=2)
            kwargs["in_period"] = True
        elif mode == "local-midnight-marked-in-period":
            snapshot.timezone = "Europe/Vienna"
            start = datetime(2026, 1, 4, 23, 30, tzinfo=UTC)
            end = start + timedelta(hours=2)
        kwargs["segments"] = [Interval(start=start, end=end)]
        if mode == "split-reversed-marked-in-period":
            # Deliberately reversed: the earliest segment defines the start day.
            kwargs["segments"].insert(0, Interval(
                start=datetime(2026, 1, 6, 0, tzinfo=UTC),
                end=datetime(2026, 1, 6, 1, tzinfo=UTC),
            ))
    snapshot.boundary_work = [BoundaryWork(
        id=" period-oracle\u00a0 ", employee_id="candidate", paid_minutes=120, **kwargs,
    )]
    return snapshot


def assert_exported_paid(snapshot, result, tmp_path, expected):
    csv_path, xlsx_path = tmp_path / "metrics.csv", tmp_path / "metrics.xlsx"
    export_table(snapshot, result, csv_path)
    export_table(snapshot, result, xlsx_path)
    with csv_path.open(encoding="utf-8-sig", newline="") as stream:
        totals = next(row for row in csv.reader(stream) if len(row) == 4 and row[0] == "candidate")
    assert int(totals[2]) == expected
    book = load_workbook(xlsx_path)
    try:
        assert book["Stundenübersicht"]["C5"].value == expected / 60
    finally:
        book.close()


@pytest.mark.parametrize("mode,expected_paid", PERIOD_CASES)
@pytest.mark.parametrize("partial", [False, True], ids=["full", "partial"])
def test_real_solver_metrics_match_start_day_export_oracle(tmp_path, mode, expected_paid, partial):
    snapshot = accounting_snapshot(mode)
    assert input_diagnostics(snapshot) == []
    assert validate(snapshot, []).complete
    result = solve(snapshot, workers=1, time_limit=5, partial=partial)
    assert result.validation.valid and result.validation.complete
    assert_exported_paid(snapshot, result, tmp_path, expected_paid)
    reported = result.metrics["employees"]["candidate"]
    assert reported["paid_minutes"] == expected_paid
    assert reported["reachable_minutes"] == expected_paid
    assert reported["deviation_minutes"] == expected_paid - snapshot.employees[0].target_minutes
    assert result.metrics["objective_contributions"]["hours"] == 0


@pytest.mark.parametrize("mode,expected_paid", PERIOD_CASES)
def test_certified_warm_fallback_uses_start_day_paid_minutes(monkeypatch, tmp_path, mode, expected_paid):
    from ortools.sat.python import cp_model

    snapshot = accounting_snapshot(mode, staffing=True)
    snapshot.assignments = [Assignment(employee_id="candidate", demand_id=snapshot.demands[0].id)]
    total = 480 + expected_paid
    snapshot.employees[0].max_period_minutes = total
    assert validate(snapshot, snapshot.assignments).complete
    original = cp_model.CpSolver.solve
    calls = []

    def certificate_then_unknown(self, model, *args, **kwargs):
        certificate = bool(self.parameters.fix_variables_to_their_hinted_value)
        calls.append(certificate)
        # Only termination of the later search is injected. The certificate,
        # full model (including the personal cap), and validator remain real.
        return original(self, model, *args, **kwargs) if certificate else cp_model.UNKNOWN

    monkeypatch.setattr(cp_model.CpSolver, "solve", certificate_then_unknown)
    result = solve(snapshot, workers=1, time_limit=5, partial=True)
    assert calls == [True, False]
    assert result.parameters["warm_start_certificate_status"] == "OPTIMAL"
    assert result.parameters["last_optimization_status"] == "UNKNOWN"
    assert result.validation.valid and result.validation.complete
    assert_exported_paid(snapshot, result, tmp_path, total)
    reported = result.metrics["employees"]["candidate"]
    assert reported["paid_minutes"] == reported["reachable_minutes"] == total
    assert reported["deviation_minutes"] == total - snapshot.employees[0].target_minutes
    assert result.metrics["objective_contributions"]["hours"] == 0
    assert result.metrics["hours_attainment"]["median"] == round(100 * total / snapshot.employees[0].target_minutes, 1)


@pytest.mark.parametrize("mode,expected_paid", PERIOD_CASES)
def test_hours_attainment_includes_only_period_personal_work(mode, expected_paid):
    snapshot = accounting_snapshot(mode, staffing=True)
    result = solve(snapshot, workers=1, time_limit=5)
    assert result.validation.valid and result.validation.complete
    assert result.metrics["hours_attainment"]["people"] == 1
    expected = round(100 * (480 + expected_paid) / snapshot.employees[0].target_minutes, 1)
    assert result.metrics["hours_attainment"]["median"] == expected


@pytest.mark.parametrize("mode,expected_paid,fixed_context", [
    (mode, paid, False) for mode, paid in PERIOD_CASES
] + [(mode, 0, True) for mode in ("prior-marked-in-period", "future-marked-in-period")])
def test_greedy_warm_load_counts_only_start_day_personal_work(monkeypatch, mode, expected_paid, fixed_context):
    from ortools.sat.python import cp_model

    snapshot = accounting_snapshot(mode, staffing=True)
    period_demand = snapshot.demands[0].id
    if fixed_context:
        work, = snapshot.boundary_work
        snapshot.boundary_work = []
        snapshot.shifts.append(Shift(
            id="fixed-context", name="Synthetic fixed context", team_id="review-team",
            kind="day", segments=work.segments, paid_minutes=work.paid_minutes,
        ))
        snapshot.demands.append(Demand(
            id="fixed-context", shift_id="fixed-context", position_id=snapshot.positions[0].id,
            minimum=1, maximum=1,
        ))
        snapshot.assignments = [Assignment(employee_id="candidate", demand_id="fixed-context", fixed=True)]
    candidate = snapshot.employees[0]
    snapshot.employees += [candidate.model_copy(deep=True, update={"id": "relief"})]
    # Exercise the large-team initial-plan branch without unrelated choices.
    snapshot.employees += [candidate.model_copy(deep=True, update={
        "id": f"excluded-{i}", "excluded": True,
    }) for i in range(38)]
    assert input_diagnostics(snapshot) == []
    original = cp_model.CpSolver.solve
    calls = []

    def certificate_then_unknown(self, model, *args, **kwargs):
        certificate = bool(self.parameters.fix_variables_to_their_hinted_value)
        calls.append(certificate)
        return original(self, model, *args, **kwargs) if certificate else cp_model.UNKNOWN

    monkeypatch.setattr(cp_model.CpSolver, "solve", certificate_then_unknown)
    result = solve(snapshot, workers=1, time_limit=5)
    assert calls == [True, False]
    assert result.parameters["warm_start_certificate_status"] == "OPTIMAL"
    assert result.validation.valid and result.validation.complete
    assert len(result.assignments) == 1 + int(fixed_context)
    # Stable ties choose the first employee. Period work breaks the tie; prior
    # and future context must not, even when the import marker says in-period.
    selected = next(a for a in result.assignments if a.demand_id == period_demand)
    assert selected.employee_id == ("relief" if expected_paid else "candidate")


@pytest.mark.parametrize("mode,expected_paid", PERIOD_CASES)
@pytest.mark.parametrize("at_cap", [False, True], ids=["one-minute-over-cap", "at-cap"])
def test_personal_paid_cap_matches_start_day_oracle(mode, expected_paid, at_cap):
    snapshot = accounting_snapshot(mode, staffing=True)
    assignment = Assignment(employee_id="candidate", demand_id=snapshot.demands[0].id)
    snapshot.assignments = [assignment.model_copy(update={"fixed": True})]
    snapshot.employees[0].max_period_minutes = 480 + expected_paid - int(not at_cap)
    manual = validate(snapshot, snapshot.assignments)
    assert manual.valid is at_cap
    if not at_cap:
        assert {d.code for d in manual.diagnostics} == {"personal_period_limit"}
    result = solve(snapshot, workers=1, time_limit=5)
    assert result.validation.valid is at_cap
    if at_cap:
        assert result.validation.complete
        assert result.metrics["employees"]["candidate"]["paid_minutes"] == 480 + expected_paid
    else:
        assert result.solver_status == "INFEASIBLE"
        assert not result.assignments


@pytest.mark.parametrize("reverse_segments", [False, True])
def test_independent_period_oracle_preserves_exact_identifiers(tmp_path, reverse_segments):
    snapshot = snapshot_for_review(days=2)
    employee = snapshot.employees[1]
    employee.id = " Candidate\u00a0ID "
    employee.credit_minutes, employee.balance_minutes = 75, -42
    snapshot.shifts, snapshot.demands = [], []
    snapshot.timezone = "Europe/Vienna"
    snapshot.profiles[0].min_rest_minutes = 0
    values = [
        (" prior ", "2026-01-04T20:00+00:00", "2026-01-04T21:00+00:00", 500),
        ("\u03bb-personal ", "2026-01-04T23:30+00:00", "2026-01-05T03:00+00:00", 300),
        (" future ", "2026-01-07T07:00+00:00", "2026-01-07T09:00+00:00", 200),
    ]
    snapshot.boundary_work = [BoundaryWork(
        id=wid, employee_id=employee.id, kind="day", paid_minutes=paid, in_period=True,
        segments=[Interval(start=datetime.fromisoformat(start), end=datetime.fromisoformat(end))],
    ) for wid, start, end, paid in values]
    split = BoundaryWork(
        id="split-preserve-id ", employee_id=employee.id, kind="night", paid_minutes=120, in_period=True,
        segments=[Interval(start=datetime.fromisoformat(a), end=datetime.fromisoformat(b)) for a, b in [
            ("2026-01-05T22:00+00:00", "2026-01-06T00:00+00:00"),
            ("2026-01-06T01:00+00:00", "2026-01-06T02:00+00:00"),
        ]],
    )
    if reverse_segments:
        split.segments.reverse()
    snapshot.boundary_work.append(split)
    snapshot.metadata["provenance"] = {split.id: {
        "name": "=literal synthetic name", "function_id": " Function\u00a0ID ",
        "workplace_id": "000301",
    }}
    result = checked_result(snapshot, [])
    csv_path, xlsx_path = tmp_path / "oracle.csv", tmp_path / "oracle.xlsx"
    export_table(snapshot, result, csv_path)
    export_table(snapshot, result, xlsx_path)
    with csv_path.open(encoding="utf-8-sig", newline="") as stream:
        csv_rows = list(csv.reader(stream))
    totals = next(row for row in csv_rows if len(row) == 4 and row[0] == employee.id)
    # Only local Jan 5 starts receive paid credit, counted once per source work.
    expected_paid = 300 + 120
    expected_actual = expected_paid + 75
    expected_balance = expected_actual - 42 - 1000
    assert tuple(map(int, totals[1:])) == (1000, expected_actual, expected_balance)
    detail = [row for row in csv_rows if len(row) == 12 and row[0] == employee.id]
    assert {row[2] for row in detail} == {w.id for w in snapshot.boundary_work}
    split_rows = [row for row in detail if row[2] == split.id]
    assert [row[8] for row in split_rows] == ["120", ""]
    assert all(row[3:6] == [" Function\u00a0ID ", "000301", "'=literal synthetic name"] for row in split_rows)
    book = load_workbook(xlsx_path)
    try:
        xlsx_rows = [["" if value is None else str(value) for value in row]
                     for row in book["Einteilungen"].values
                     if row[0] == employee.id and row[2] in {w.id for w in snapshot.boundary_work}]
        assert xlsx_rows == detail
        hours = book["Stundenübersicht"]
        assert hours["C6"].value == expected_paid / 60
        assert hours["D6"].value == 75 / 60
        assert hours["F6"].value == expected_balance / 60
        assert hours["G6"].value == 2
        assert all(cell.data_type != "f" for sheet in book for row in sheet for cell in row)
    finally:
        book.close()


@pytest.mark.parametrize("exists", [False, True], ids=["new", "existing"])
def test_csv_publish_failure_preserves_destination_and_removes_temporary(tmp_path, monkeypatch, exists):
    from pathlib import Path

    snapshot = snapshot_for_review(days=1)
    result = checked_result(snapshot, [])
    target = tmp_path / "synthetic.csv"
    if exists:
        target.write_bytes(b"previous synthetic export")
    staged = []

    def fail_replace(source, destination):
        assert destination == target
        assert source.parent == target.parent
        assert source.read_bytes().startswith(b"\xef\xbb\xbfZeitraum,")
        staged.append(source)
        raise OSError("synthetic publish failure")

    monkeypatch.setattr(Path, "replace", fail_replace)
    with pytest.raises(OSError, match="synthetic publish failure"):
        export_table(snapshot, result, target)
    assert len(staged) == 1 and not staged[0].exists()
    if exists:
        assert target.read_bytes() == b"previous synthetic export"
    else:
        assert not target.exists()
    assert set(tmp_path.iterdir()) == ({target} if exists else set())
