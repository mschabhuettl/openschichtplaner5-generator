from sp5generator.export import safe_cell
import csv

import pytest


def test_spreadsheet_formulas_are_escaped():
    for text in ["=1+1", " +2", "-3", "@SUM(A1:A2)", "\t=1", "\r=1"]:
        assert safe_cell(text).startswith("'")
    assert safe_cell("Testperson 001") == "Testperson 001"
    assert safe_cell(-23) == -23


def test_context_duty_does_not_change_period_balance():
    from datetime import timedelta
    from sp5generator.demo import make_demo
    from sp5generator.models import Assignment, Result, Validation
    from sp5generator.export import rows

    snapshot = make_demo(days=1)
    shift = snapshot.shifts[0]
    shift.segments = [
        i.model_copy(
            update={
                "start": i.start - timedelta(days=2),
                "end": i.end - timedelta(days=2),
            }
        )
        for i in shift.segments
    ]
    result = Result(
        snapshot_id=snapshot.id,
        snapshot_hash="unused",
        solver_status="FEASIBLE",
        assignments=[Assignment(employee_id="e000", demand_id="d0-day-p0")],
        validation=Validation(valid=True, complete=False, diagnostics=[]),
        runtime_seconds=0,
    )
    output = list(rows(snapshot, result))
    balance = next(row for row in output if len(row) == 4 and row[0] == "e000")
    assert balance[2] == snapshot.employees[0].credit_minutes


def incomplete_result(snapshot):
    from sp5generator.domain import snapshot_hash
    from sp5generator.models import Result, Validation

    return Result(
        snapshot_id=snapshot.id,
        snapshot_hash=snapshot_hash(snapshot),
        solver_status="FEASIBLE",
        vacancies={"stale-demand": 99},
        validation=Validation(valid=True, complete=True, diagnostics=[]),
        runtime_seconds=0,
    )


@pytest.mark.parametrize("text", ["#REF!", "#N/A", "#VALUE!"])
def test_xlsx_error_like_names_remain_literal_text(tmp_path, text):
    from openpyxl import load_workbook
    from sp5generator.demo import make_demo
    from sp5generator.export import export_table
    from sp5generator.models import Assignment

    snapshot = make_demo(days=1)
    snapshot.assignments = []
    snapshot.employees[2].name = text
    result = incomplete_result(snapshot)
    result.assignments = [Assignment(employee_id=snapshot.employees[2].id,
                                     demand_id="d0-day-p0")]
    path = tmp_path / "literal.xlsx"
    export_table(snapshot, result, path)
    workbook = load_workbook(path)
    try:
        cells = [cell for sheet in workbook for row in sheet for cell in row
                 if cell.value == text]
        assert len(cells) == 3  # Calendar, balances and assignment detail.
        assert all(cell.data_type == "s" for cell in cells)
        assert workbook["Stundenübersicht"]["C7"].data_type == "n"
    finally:
        workbook.close()


@pytest.mark.parametrize("text", ["x" * 32768, "private-name\x01", "private-name\uffff"],
                         ids=["overlong", "control", "invalid-xml"])
def test_xlsx_unsupported_text_is_rejected_without_replacing_export(tmp_path, text):
    from sp5generator.demo import make_demo
    from sp5generator.export import export_table

    snapshot = make_demo(days=1)
    snapshot.assignments = []
    snapshot.employees[0].name = text
    result = incomplete_result(snapshot)
    path = tmp_path / "previous.xlsx"
    path.write_bytes(b"previous export")
    with pytest.raises(ValueError, match="Excel") as error:
        export_table(snapshot, result, path)
    assert "private-name" not in str(error.value)
    assert path.read_bytes() == b"previous export"


def test_xlsx_supported_text_is_preserved_at_cell_length_boundary(tmp_path):
    from openpyxl import load_workbook
    from sp5generator.demo import make_demo
    from sp5generator.export import export_table

    snapshot = make_demo(days=1)
    snapshot.assignments = []
    text = "Ä\t\n" + "x" * (32767 - 3)
    snapshot.employees[0].name = text
    path = tmp_path / "boundary.xlsx"
    export_table(snapshot, incomplete_result(snapshot), path)
    workbook = load_workbook(path)
    try:
        assert workbook.active["A5"].value == text
        assert workbook["Stundenübersicht"]["A5"].value == text
    finally:
        workbook.close()


@pytest.mark.parametrize("suffix", ["csv", "xlsx"])
def test_partial_export_revalidates_and_recomputes_vacancies(tmp_path, suffix):
    from sp5generator.demo import make_demo
    from sp5generator.export import export_table

    snapshot = make_demo(days=1)
    snapshot.assignments = []
    result = incomplete_result(snapshot)
    target = tmp_path / f"partial.{suffix}"
    export_table(snapshot, result, target)
    if suffix == "csv":
        with target.open(encoding="utf-8-sig", newline="") as stream:
            output = list(csv.reader(stream))
    else:
        from openpyxl import load_workbook
        workbook = load_workbook(target)
        output = list(workbook["Einteilungen"].values)
        workbook.close()
    assert output[1][2] == "Nicht vollständig"
    assert any(row[0] == snapshot.demands[0].id and str(row[1]) == str(snapshot.demands[0].minimum)
               for row in output if row)
    assert any(row[0] == "vacancy" for row in output if row)
    assert all(row[0] != "stale-demand" for row in output if row)


@pytest.mark.parametrize("invalid", ["unknown-person", "duplicate", "interval-mismatch"])
def test_invalid_export_leaves_existing_file_unchanged(tmp_path, invalid):
    from datetime import timedelta
    from sp5generator.demo import make_demo
    from sp5generator.export import export_table
    from sp5generator.models import Assignment

    snapshot = make_demo(days=1)
    result = incomplete_result(snapshot)
    assignment = Assignment(employee_id="e000", demand_id=snapshot.demands[0].id)
    if invalid == "unknown-person":
        assignment.employee_id = "missing"
    if invalid == "interval-mismatch":
        assignment.segments = [segment.model_copy(update={"end": segment.end + timedelta(minutes=1)})
                               for segment in snapshot.shifts[0].segments]
    result.assignments = [assignment, assignment] if invalid == "duplicate" else [assignment]
    target = tmp_path / "plan.csv"
    target.write_bytes(b"previous export")
    with pytest.raises(ValueError, match="Ungültigen Entwurf"):
        export_table(snapshot, result, target)
    assert target.read_bytes() == b"previous export"


def test_export_checks_snapshot_identity(tmp_path):
    from sp5generator.demo import make_demo
    from sp5generator.export import export_table

    snapshot = make_demo(days=1)
    result = incomplete_result(snapshot)
    result.snapshot_id = "another-snapshot"
    with pytest.raises(ValueError, match="reference"):
        export_table(snapshot, result, tmp_path / "plan.csv")


def test_workbook_calendar_balances_and_untrusted_names(tmp_path):
    from openpyxl import load_workbook
    from sp5generator.demo import make_demo
    from sp5generator.domain import snapshot_hash
    from sp5generator.export import export_table
    from sp5generator.models import Assignment

    snapshot = make_demo(days=2)
    snapshot.assignments = []
    snapshot.metadata["project_name"] = "=untrusted project"
    employee = snapshot.employees[2]
    employee.name = "=untrusted person"
    employee.credit_minutes = 120
    employee.balance_minutes = -60
    result = incomplete_result(snapshot)
    result.snapshot_hash = snapshot_hash(snapshot)
    result.assignments = [Assignment(employee_id=employee.id, demand_id="d0-night-p0", fixed=True)]
    path = tmp_path / "calendar.xlsx"
    export_table(snapshot, result, path)
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Plan 2026-01", "Stundenübersicht", "Offene Stellen", "Einteilungen"]
    calendar = workbook.active
    assert calendar.freeze_panes == "B5"
    assert calendar["A1"].data_type == "s"
    assert calendar["A7"].value == "'=untrusted person"
    assert "20:00–06:00 (+1)" in calendar["B7"].value
    assert "20:00–06:00 (+1)" in calendar["C7"].value
    assert calendar["B7"].fill.fgColor.rgb.endswith("EAEFFB")
    balances = workbook["Stundenübersicht"]
    assert balances["C7"].value == 10  # Midnight crossing is paid once.
    assert balances["D7"].value == 2
    assert balances["E7"].value == -1
    assert balances["F7"].value == 11
    assert balances["G7"].value == 2
    assert balances["H7"].value == 1
    assert all(cell.data_type != "f" for sheet in workbook for row in sheet for cell in row)
    workbook.close()
