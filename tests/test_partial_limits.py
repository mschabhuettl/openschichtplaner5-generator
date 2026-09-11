"""Synthetic counterexamples for the reported 0.9.29 partial-plan symptoms.

No customer snapshot is used. Expected elapsed totals are stated explicitly,
not obtained from the solver's daily-minute helper.
"""

from datetime import UTC, date, datetime, timedelta

import pytest
from ortools.sat.python import cp_model

from sp5generator import solver
from sp5generator.models import Assignment, Interval
from sp5generator.timeutils import localize
from sp5generator.validator import validate
from test_core_rules import case, shift


def plan(snapshot):
    return [Assignment(employee_id="e0", demand_id=d.id) for d in snapshot.demands]


@pytest.mark.parametrize("partial", [False, True])
@pytest.mark.parametrize("limit,expected_count", [(479, 0), (480, 1), (960, 2)])
def test_daily_limit_sums_separate_duties_not_paid_minutes(partial, limit, expected_count):
    snapshot = case(1, [shift("morning", 5, 0, 8), shift("evening", 5, 16, 8)])
    snapshot.profiles[0].min_rest_minutes = 0
    snapshot.profiles[0].max_daily_minutes = limit
    # Each duty pays just one hour; the hard cap counts eight elapsed hours.
    for duty in snapshot.shifts:
        duty.paid_minutes = 60
    checked = validate(snapshot, plan(snapshot))
    assert checked.valid is (expected_count == 2)
    if expected_count < 2:
        assert "daily_limit" in {d.code for d in checked.diagnostics}
    result = solver.solve(snapshot, 3, partial=partial)
    if not partial and expected_count < 2:
        assert result.solver_status == "INFEASIBLE"
    else:
        assert result.solver_status == "OPTIMAL"
        assert len(result.assignments) == expected_count
        assert validate(snapshot, result.assignments).valid
        assert sum(result.vacancies.values()) == 2 - expected_count


@pytest.mark.parametrize("partial", [False, True])
def test_weekly_limit_uses_fixed_context_and_not_target_or_paid_minutes(partial):
    snapshot = case(1, [shift("boundary", 5, 8, 8), shift("new", 7, 8, 8)])
    snapshot.period_start = date(2026, 1, 7)
    snapshot.profiles[0].max_weekly_minutes = 15 * 60
    snapshot.employees[0].target_minutes = 100 * 60
    snapshot.shifts[0].paid_minutes = snapshot.shifts[1].paid_minutes = 60
    snapshot.assignments = [Assignment(employee_id="e0", demand_id="boundary", fixed=True)]
    checked = validate(snapshot, plan(snapshot))
    assert not checked.valid
    assert "weekly_limit" in {d.code for d in checked.diagnostics}
    result = solver.solve(snapshot, 3, partial=partial)
    if partial:
        assert result.solver_status == "OPTIMAL"
        assert [(a.demand_id, a.fixed) for a in result.assignments] == [("boundary", True)]
        assert result.vacancies == {"new": 1}
        assert validate(snapshot, result.assignments).valid
    else:
        assert result.solver_status == "INFEASIBLE"


@pytest.mark.parametrize("maximum,expected", [(None, 3), (40 * 60, 1)])
def test_24_hour_duties_are_not_forbidden_by_11_36_rest_alone(maximum, expected):
    snapshot = case(1, [shift("a", 5, 0, 24), shift("b", 7, 0, 24), shift("c", 9, 0, 24)])
    profile = snapshot.profiles[0]
    profile.min_rest_minutes = 660
    profile.weekly_rest_minutes = 2160
    profile.max_weekly_minutes = maximum
    snapshot.employees[0].target_minutes = 40 * 60
    # 72h elapsed, 24h paid, 40h target. All three quantities differ.
    for duty in snapshot.shifts:
        duty.paid_minutes = 8 * 60
    result = solver.solve(snapshot, 3, partial=True)
    assert result.solver_status == "OPTIMAL"
    assert len(result.assignments) == expected
    assert validate(snapshot, result.assignments).valid
    assert result.metrics["employees"]["e0"]["paid_minutes"] == expected * 480


@pytest.mark.parametrize("daily_limit,accepted", [(719, False), (720, True)])
def test_daily_limit_is_not_a_single_duty_length_limit(daily_limit, accepted):
    snapshot = case(1, [shift("noon_to_noon", 5, 12, 24)])
    snapshot.profiles[0].min_rest_minutes = 660
    snapshot.profiles[0].weekly_rest_minutes = 2160
    snapshot.profiles[0].max_daily_minutes = daily_limit
    # Noon to noon is 12h on each local calendar day, not 24h on one day.
    assert validate(snapshot, plan(snapshot)).valid is accepted
    result = solver.solve(snapshot, 3, partial=True)
    assert len(result.assignments) == int(accepted)
    assert result.validation.valid


@pytest.mark.parametrize("partial", [False, True])
@pytest.mark.parametrize("limit,code,day,hour,duration,cap", [
    ("max_daily_minutes", "daily_limit", 5, 20, 20, 720),
    ("max_weekly_minutes", "weekly_limit", 11, 23, 9, 420),
])
def test_hard_limits_cover_spill_after_period_end(partial, limit, code, day, hour, duration, cap):
    snapshot = case(1, [shift("spill", day, hour, duration)])
    snapshot.period_start = snapshot.period_end = date(2026, 1, day)
    setattr(snapshot.profiles[0], limit, cap)
    snapshot.shifts[0].paid_minutes = 60
    # Monday 20:00 -> Tuesday 16:00: 16h exceed the 12h daily cap.
    # Sunday 23:00 -> Monday 08:00: 8h exceed the 7h next-ISO-week cap.
    checked = validate(snapshot, plan(snapshot))
    assert not checked.valid
    assert code in {d.code for d in checked.diagnostics}
    result = solver.solve(snapshot, 3, partial=partial)
    assert not result.assignments
    assert result.solver_status == ("OPTIMAL" if partial else "INFEASIBLE")
    if partial:
        assert result.validation.valid
        assert result.vacancies == {"spill": 1}


@pytest.mark.parametrize("limit,code", [
    ("max_daily_minutes", "daily_limit"), ("max_weekly_minutes", "weekly_limit"),
])
@pytest.mark.parametrize("cap,accepted", [(599, False), (600, True)])
def test_spill_limits_include_fixed_following_context(limit, code, cap, accepted):
    snapshot = case(1, [shift("spill", 11, 23, 9), shift("fixed", 12, 12, 2)])
    snapshot.period_start = snapshot.period_end = date(2026, 1, 11)
    profile = snapshot.profiles[0]
    profile.min_rest_minutes = 0
    setattr(profile, limit, cap)
    snapshot.assignments = [Assignment(employee_id="e0", demand_id="fixed", fixed=True)]
    # The next local day and ISO week contain 8h of spill + 2h fixed work.
    checked = validate(snapshot, plan(snapshot))
    assert checked.valid is accepted
    if not accepted:
        assert code in {d.code for d in checked.diagnostics}
    result = solver.solve(snapshot, 3, partial=True)
    assert result.solver_status == "OPTIMAL"
    assert {a.demand_id for a in result.assignments} == ({"spill", "fixed"} if accepted else {"fixed"})
    assert result.validation.valid
    assert validate(snapshot, result.assignments).valid


@pytest.mark.parametrize("limit", ["max_daily_minutes", "max_weekly_minutes"])
def test_unselected_spill_does_not_activate_unrelated_future_context_limit(limit):
    snapshot = case(1, [shift("spill", 11, 23, 9), shift("fixed", 12, 12, 10)])
    snapshot.period_start = snapshot.period_end = date(2026, 1, 11)
    snapshot.profiles[0].min_rest_minutes = 0
    setattr(snapshot.profiles[0], limit, 9 * 60)
    snapshot.assignments = [Assignment(employee_id="e0", demand_id="fixed", fixed=True)]
    # Fixed future work alone exceeds the cap, outside this planning period.
    # Do not reject all partial plans merely because a spill candidate exists.
    checked = validate(snapshot, snapshot.assignments)
    assert checked.valid and not checked.complete
    result = solver.solve(snapshot, 3, partial=True)
    assert result.solver_status == "OPTIMAL"
    assert {a.demand_id for a in result.assignments} == {"fixed"}
    assert result.validation.valid


@pytest.mark.parametrize("limit", ["max_daily_minutes", "max_weekly_minutes"])
@pytest.mark.parametrize("assigned,active", [(True, True), (False, True), (True, False)])
def test_spill_uses_assigned_profile_valid_on_tail_date(limit, assigned, active):
    snapshot = case(1, [shift("spill", 11, 23, 9)])
    snapshot.period_start = snapshot.period_end = date(2026, 1, 11)
    tail_profile = snapshot.profiles[0].model_copy(update={
        "id": "tail", "valid_from": date(2026, 1, 12 if active else 13), limit: 7 * 60,
    })
    snapshot.profiles.append(tail_profile)
    if assigned:
        snapshot.employees[0].profile_ids.append("tail")
    accepted = not (assigned and active)
    assert validate(snapshot, plan(snapshot)).valid is accepted
    result = solver.solve(snapshot, 3, partial=True)
    assert len(result.assignments) == int(accepted)
    assert result.validation.valid


def test_spill_does_not_extend_period_minutes_or_paid_target_accounting():
    snapshot = case(1, [shift("spill", 11, 23, 9)])
    snapshot.period_start = snapshot.period_end = date(2026, 1, 11)
    snapshot.profiles[0].max_period_minutes = 60
    snapshot.profiles[0].max_daily_minutes = 480
    snapshot.profiles[0].max_weekly_minutes = 480
    snapshot.shifts[0].paid_minutes = 120
    result = solver.solve(snapshot, 3)
    assert result.solver_status == "OPTIMAL" and result.validation.complete
    assert result.metrics["employees"]["e0"]["paid_minutes"] == 120
    assert validate(snapshot, result.assignments).complete


def test_long_spill_needs_context_to_end_of_its_last_iso_week():
    # A deliberately extreme, but contract-supported synthetic duty. There is
    # no invented duration cap or weekly-rest rule in this fixture.
    snapshot = case(1, [shift("spill", 4, 23, 177)])
    snapshot.period_start = snapshot.period_end = date(2026, 1, 4)
    snapshot.context_end = date(2026, 1, 12)
    snapshot.profiles[0].max_weekly_minutes = 20000
    checked = validate(snapshot, plan(snapshot))
    assert checked.valid and not checked.complete
    assert "context" in {d.code for d in checked.diagnostics}
    result = solver.solve(snapshot, 3)
    assert result.solver_status == "OPTIMAL"
    assert result.validation.valid and not result.validation.complete
    snapshot.context_end = date(2026, 1, 18)
    assert validate(snapshot, plan(snapshot)).complete


@pytest.mark.parametrize("kind", ["overlap", "rest", "same_shift"])
def test_partial_never_keeps_conflicting_assignments(kind):
    snapshot = case(1, [shift("a", 5, 8, 8), shift("b", 5, 16, 8)])
    snapshot.profiles[0].min_rest_minutes = 660
    if kind == "overlap":
        snapshot.shifts[1] = shift("b", 5, 12, 8)
    elif kind == "same_shift":
        snapshot.demands[1].shift_id = "a"
    assert not validate(snapshot, plan(snapshot)).valid
    result = solver.solve(snapshot, 3, partial=True)
    assert len(result.assignments) == 1
    assert validate(snapshot, result.assignments).valid


def dated_case(start, end, intervals):
    snapshot = case(1, [shift(str(i), 5, 0, 1) for i in range(len(intervals))])
    snapshot.timezone = "Europe/Vienna"
    snapshot.period_start, snapshot.period_end = start, end
    snapshot.context_start, snapshot.context_end = start - timedelta(days=14), end + timedelta(days=14)
    profile = snapshot.profiles[0]
    profile.valid_from, profile.valid_until = snapshot.context_start, snapshot.context_end
    employee = snapshot.employees[0]
    employee.employment_start, employee.employment_end = profile.valid_from, profile.valid_until
    employee.approvals[0].valid_from, employee.approvals[0].valid_until = profile.valid_from, profile.valid_until
    for duty, (a, b) in zip(snapshot.shifts, intervals):
        duty.segments = [Interval(start=a, end=b)]
        duty.paid_minutes = 60
    return snapshot


@pytest.mark.parametrize("daily_cap,accepted", [(240, False), (300, True)])
def test_autumn_dst_counts_five_elapsed_hours_in_four_wall_hours(daily_cap, accepted):
    day = date(2026, 10, 25)
    snapshot = dated_case(day, day, [(localize(day, "00:00", "Europe/Vienna"), localize(day, "04:00", "Europe/Vienna"))])
    snapshot.profiles[0].max_daily_minutes = daily_cap
    assert validate(snapshot, plan(snapshot)).valid is accepted
    result = solver.solve(snapshot, 3, partial=True)
    assert len(result.assignments) == int(accepted)
    assert result.validation.valid


def test_week_crossing_duty_is_split_at_local_monday_and_iso_year():
    # Sunday 23:00 -> Monday 03:00 local: one hour in 2025-W52,
    # three hours in 2026-W01. Wednesday adds two in that same ISO week.
    start, end = date(2025, 12, 28), date(2025, 12, 31)
    snapshot = dated_case(start, end, [
        (datetime(2025, 12, 28, 22, tzinfo=UTC), datetime(2025, 12, 29, 2, tzinfo=UTC)),
        (datetime(2025, 12, 31, 7, tzinfo=UTC), datetime(2025, 12, 31, 9, tzinfo=UTC)),
    ])
    profile = snapshot.profiles[0]
    profile.max_weekly_minutes = 4 * 60
    checked = validate(snapshot, plan(snapshot))
    assert not checked.valid
    assert any(d.code == "weekly_limit" and d.date == "2025-12-29" for d in checked.diagnostics)
    assert len(solver.solve(snapshot, 3, partial=True).assignments) == 1
    profile.max_weekly_minutes = 5 * 60
    assert validate(snapshot, plan(snapshot)).valid
    assert len(solver.solve(snapshot, 3, partial=True).assignments) == 2


def test_strictest_assigned_profile_wins_but_unassigned_limit_does_not_apply():
    snapshot = case(1, [shift("a", 5, 8, 8), shift("b", 6, 8, 8)])
    strict = snapshot.profiles[0].model_copy(update={"id": "strict", "max_weekly_minutes": 480})
    snapshot.profiles.append(strict)
    assert len(solver.solve(snapshot, 3, partial=True).assignments) == 2
    snapshot.employees[0].profile_ids.append("strict")
    assert not validate(snapshot, plan(snapshot)).valid
    result = solver.solve(snapshot, 3, partial=True)
    assert len(result.assignments) == 1
    assert result.validation.valid


@pytest.mark.parametrize("credit,balance", [(480, 0), (0, 480), (720, -240)])
def test_quality_timeout_preserves_independently_validated_partial_incumbent(monkeypatch, credit, balance):
    snapshot = case(1, [shift("a", 5, 8, 8), shift("b", 6, 8, 8)])
    snapshot.profiles[0].max_weekly_minutes = 480
    snapshot.employees[0].target_minutes = 480
    snapshot.employees[0].credit_minutes = credit
    snapshot.employees[0].balance_minutes = balance
    for duty in snapshot.shifts:
        duty.paid_minutes = 60
    original = cp_model.CpSolver.solve
    calls = []

    def first_solution_then_unknown(self, model, *args, **kwargs):
        calls.append(1)
        if len(calls) == 2:
            return cp_model.UNKNOWN
        return original(self, model, *args, **kwargs)

    monkeypatch.setattr(cp_model.CpSolver, "solve", first_solution_then_unknown)
    result = solver.solve(snapshot, 3, partial=True)
    assert len(calls) == 2
    assert result.solver_status == "FEASIBLE"
    assert result.parameters["last_optimization_status"] == "UNKNOWN"
    assert len(result.assignments) == 1
    assert result.validation.valid and not result.validation.complete
    assert validate(snapshot, result.assignments).valid
    metrics = result.metrics["employees"]["e0"]
    assert metrics["paid_minutes"] == 60
    assert metrics["credit_minutes"] == credit
    assert metrics["balance_minutes"] == balance
    assert metrics["deviation_minutes"] == 60
    assert result.metrics["objective_contributions"]["hours"] == 60
    assert result.metrics["objective_phase"] == "vacancies"
    assert result.objective_value == 1  # Vacancy count, not the hours-quality cost.


def test_feasible_partial_can_finish_before_hours_and_block_optimization(monkeypatch):
    snapshot = case(1, [shift("a", 5, 8, 8), shift("b", 6, 8, 8)])
    snapshot.profiles[0].max_weekly_minutes = 480
    original = cp_model.CpSolver.solve
    calls = []

    def feasible_without_optimality_proof(self, model, *args, **kwargs):
        calls.append(1)
        assert original(self, model, *args, **kwargs) == cp_model.OPTIMAL
        return cp_model.FEASIBLE  # Exercise the time-limited first-phase branch.

    monkeypatch.setattr(cp_model.CpSolver, "solve", feasible_without_optimality_proof)
    result = solver.solve(snapshot, 3, partial=True)
    assert len(calls) == 1
    assert result.solver_status == "FEASIBLE"
    assert result.metrics["objective_phase"] == "vacancies"
    assert len(result.assignments) == 1
    assert result.validation.valid and not result.validation.complete
    assert validate(snapshot, result.assignments).valid


def test_unknown_without_incumbent_never_returns_unchecked_assignments(monkeypatch):
    monkeypatch.setattr(cp_model.CpSolver, "solve", lambda *args, **kwargs: cp_model.UNKNOWN)
    result = solver.solve(case(1), 3, partial=True)
    assert result.solver_status == "UNKNOWN"
    assert not result.assignments and not result.validation.valid
    assert result.metrics["planning_diagnostics"]["employees"]["e0"]["reason"] == "no_valid_plan"


@pytest.mark.parametrize("credit,balance", [(480, 0), (0, 480), (720, -240)])
@pytest.mark.parametrize("certified", [False, True])
def test_initial_plan_timeout_requires_certificate_and_keeps_account_balance(
    monkeypatch, credit, balance, certified,
):
    from sp5generator.models import Objectives

    # Forty employees activate the existing full-plan initial proposal path.
    # This is deliberately not a partial-plan or 600-second load reproduction.
    snapshot = case(40)
    snapshot.profiles[0].max_daily_minutes = 480
    snapshot.profiles[0].max_weekly_minutes = 480
    snapshot.objectives = Objectives(
        hours=1, changes=0, nights=0, weekends=0, holidays=0, wishes=0,
        workday_transitions=0,
    )
    for employee in snapshot.employees:
        employee.target_minutes = 480
        employee.credit_minutes = credit
        employee.balance_minutes = balance
    original = cp_model.CpSolver.solve
    calls = []

    def certificate_then_unknown(self, model, *args, **kwargs):
        calls.append(bool(self.parameters.fix_variables_to_their_hinted_value))
        if len(calls) == 1 and certified:
            return original(self, model, *args, **kwargs)
        return cp_model.UNKNOWN

    monkeypatch.setattr(cp_model.CpSolver, "solve", certificate_then_unknown)
    result = solver.solve(snapshot, 10)
    assert calls == [True, False]
    assert result.parameters["last_optimization_status"] == "UNKNOWN"
    if not certified:
        assert result.parameters["warm_start_certificate_status"] == "UNKNOWN"
        assert result.solver_status == "UNKNOWN"
        assert not result.assignments and not result.validation.valid
        assert all(d["reason"] == "no_valid_plan" for d in
                   result.metrics["planning_diagnostics"]["employees"].values())
        return
    assert result.parameters["warm_start_certificate_status"] == "OPTIMAL"
    assert result.solver_status == "FEASIBLE"
    assert result.metrics["objective_phase"] == "validated_initial_solution"
    assert result.validation.complete and validate(snapshot, result.assignments).complete
    assert len(result.assignments) == 1
    assigned = result.assignments[0].employee_id
    for eid, metrics in result.metrics["employees"].items():
        paid = 480 if eid == assigned else 0
        assert metrics["paid_minutes"] == paid
        assert metrics["credit_minutes"] == credit
        assert metrics["balance_minutes"] == balance
        assert metrics["deviation_minutes"] == paid
    assert result.metrics["objective_contributions"]["hours"] == 480
    assert result.metrics["weighted_objective_contributions"]["hours"] == 480
    assert result.objective_value == 480


def test_no_rule_requires_every_eligible_employee_to_receive_a_duty():
    snapshot = case(3)
    for employee in snapshot.employees:
        employee.target_minutes = 480
    result = solver.solve(snapshot, 3, partial=True)
    assert result.validation.complete
    assert len(result.assignments) == 1  # demand.maximum=1, three eligible people
    assert sum(m["paid_minutes"] == 0 for m in result.metrics["employees"].values()) == 2
    diagnostics = result.metrics["planning_diagnostics"]
    assert diagnostics["objective_weights"] == snapshot.objectives.model_dump()
    assert sorted(d["reason"] for d in diagnostics["employees"].values()) == [
        "assigned", "not_selected_with_candidates", "not_selected_with_candidates"
    ]
    assert all(d["eligible_demands"] == 1 for d in diagnostics["employees"].values())


def test_linear_hours_target_can_tie_while_block_goal_concentrates_work():
    from sp5generator.models import Objectives

    snapshot = case(3, [shift("a", 5, 8, 8), shift("b", 6, 8, 8), shift("c", 7, 8, 8)])
    for employee in snapshot.employees:
        employee.target_minutes = 40 * 60
    snapshot.profiles[0].min_rest_minutes = 660
    snapshot.profiles[0].weekly_rest_minutes = 2160
    snapshot.objectives = Objectives(
        hours=1, changes=0, nights=0, weekends=0, holidays=0, wishes=0,
        workday_transitions=100,
    )
    distributed = [
        Assignment(employee_id=f"e{i}", demand_id=d.id)
        for i, d in enumerate(snapshot.demands)
    ]
    assert validate(snapshot, distributed).complete
    result = solver.solve(snapshot, 3, partial=True)
    assert result.solver_status == "OPTIMAL" and result.validation.complete
    assert len({a.employee_id for a in result.assignments}) == 1
    costs = result.metrics["objective_contributions"]
    # All people remain below target: total L1 shortfall is identical.
    assert costs["hours"] == 3 * (2400 - 480) == 5760
    # One three-day block has two transitions; three one-day blocks have six.
    assert costs["workday_transitions"] == 2
    assert result.objective_value == 5760 + 2 * 100
    assert sum(d["reason"] == "not_selected_with_candidates" for d in result.metrics["planning_diagnostics"]["employees"].values()) == 2


@pytest.mark.parametrize("excluded,code", [
    ("approval", "approval"), ("employment", "employment"),
    ("availability", "availability"), ("team", "team"),
])
def test_nonassignment_diagnostics_use_actual_eligibility(excluded, code):
    from sp5generator.models import Availability

    snapshot = case(1)
    employee = snapshot.employees[0]
    if excluded == "approval":
        employee.approvals.clear()
    elif excluded == "employment":
        employee.employment_start = date(2026, 1, 6)
    elif excluded == "availability":
        employee.availability = [Availability(valid_from=date(2026, 1, 5), valid_until=date(2026, 1, 11), weekdays=[6])]
    else:
        employee.team_ids.clear()
    result = solver.solve(snapshot, 3, partial=True)
    diagnostic = result.metrics["planning_diagnostics"]["employees"]["e0"]
    assert diagnostic == {
        "positive_capacity_demands": 1, "eligible_demands": 0,
        "eligible_required_demands": 0, "exclusions": {code: 1},
        "assigned_demands": 0, "reason": "individually_ineligible",
    }
    assert result.validation.valid and not result.validation.complete


@pytest.mark.parametrize("partial", [False, True])
@pytest.mark.parametrize("code", ["kind", "weekend", "holiday", "qualification", "restriction", "absence"])
def test_remaining_individual_exclusions_and_explicit_recovery(partial, code):
    from sp5generator.models import Restriction

    # Saturday keeps the weekend predicate explicit, with full weekly context.
    snapshot = case(1, [shift("s", 10, 8, 8)])
    employee = snapshot.employees[0]
    duty = snapshot.shifts[0]
    if code == "kind":
        employee.allowed_kinds = ["night"]
    elif code == "weekend":
        employee.allow_weekends = False
    elif code == "holiday":
        duty.holiday = True
        employee.allow_holidays = False
    elif code == "qualification":
        snapshot.positions[0].qualifications_required = True
        snapshot.positions[0].qualification_ids = ["explicit-test-qualification"]
    elif code == "restriction":
        snapshot.restrictions = [Restriction(employee_id="e0", shift_id="s", level=1)]
    else:
        employee.unavailable = [duty.segments[0].model_copy()]

    counterplan = validate(snapshot, plan(snapshot))
    assert not counterplan.valid
    assert {d.code for d in counterplan.diagnostics} == {code}
    result = solver.solve(snapshot, 3, partial=partial)
    assert result.assignments == []
    assert result.solver_status == ("OPTIMAL" if partial else "INFEASIBLE")
    assert result.validation.valid is partial
    if partial:
        diagnostic = result.metrics["planning_diagnostics"]["employees"]["e0"]
        assert diagnostic["reason"] == "individually_ineligible"
        assert diagnostic["eligible_demands"] == 0
        assert diagnostic["exclusions"] == {code: 1}

    # Deliberate synthetic setup edits, never implicit production relaxation.
    if code == "kind":
        employee.allowed_kinds.append("day")
    elif code == "weekend":
        employee.allow_weekends = True
    elif code == "holiday":
        employee.allow_holidays = True
    elif code == "qualification":
        snapshot.positions[0].qualifications_required = False
    elif code == "restriction":
        snapshot.restrictions[0].approved = True
    else:
        # Half-open absence ending at duty start must not exclude the duty.
        employee.unavailable = [Interval(start=duty.segments[0].start - timedelta(hours=1),
                                         end=duty.segments[0].start)]
    recovered = solver.solve(snapshot, 3, partial=partial)
    assert recovered.solver_status == "OPTIMAL"
    assert len(recovered.assignments) == 1
    assert validate(snapshot, recovered.assignments).valid
    assert recovered.validation.complete
    assert recovered.metrics["planning_diagnostics"]["employees"]["e0"]["reason"] == "assigned"
    if code == "restriction":
        # An approval for an on-request restriction never overrides 'never'.
        snapshot.restrictions[0].level = 2
        forbidden = solver.solve(snapshot, 3, partial=partial)
        assert forbidden.assignments == []
        assert not validate(snapshot, plan(snapshot)).valid


@pytest.mark.parametrize("maximum,reason", [(0, "no_positive_capacity_demand"), (1, "not_selected_with_candidates")])
def test_optional_and_zero_capacity_demands_are_distinguished(maximum, reason):
    snapshot = case(1)
    snapshot.demands[0].minimum = 0
    snapshot.demands[0].maximum = maximum
    result = solver.solve(snapshot, 3, partial=True)
    diagnostic = result.metrics["planning_diagnostics"]["employees"]["e0"]
    assert diagnostic["reason"] == reason
    assert diagnostic["eligible_required_demands"] == 0
    assert diagnostic["eligible_demands"] == maximum
    assert diagnostic["exclusions"] == ({"zero_capacity": 1} if maximum == 0 else {})
    assert result.validation.complete


def test_candidate_exclusions_count_each_reason_without_claiming_disjoint_people():
    snapshot = case(1)
    snapshot.employees[0].approvals.clear()
    snapshot.employees[0].team_ids.clear()
    result = solver.solve(snapshot, 3, partial=True)
    diagnostic = result.metrics["planning_diagnostics"]["employees"]["e0"]
    assert diagnostic["positive_capacity_demands"] == 1
    assert diagnostic["exclusions"] == {"approval": 1, "team": 1}


def test_failed_input_check_does_not_publish_candidate_analysis():
    snapshot = case(1)
    snapshot.profiles[0].confirmed = False
    result = solver.solve(snapshot, 3, partial=True)
    assert result.solver_status == "MODEL_INVALID"
    assert "planning_diagnostics" not in result.metrics


@pytest.mark.parametrize("partial", [False, True])
@pytest.mark.parametrize("missing", ["approval", "team", "kind"])
def test_fixed_boundary_time_still_requires_assignment_eligibility(partial, missing):
    """Document why clearing context mapping diagnostics is not a safe fix.

    Known historical time currently uses the assignment contract, not a
    separate person/time ledger. Even a fully confirmed input cannot bypass
    its placement checks, including in partial mode.
    """
    snapshot = case(1, [shift("boundary", 5, 8, 8), shift("new", 7, 8, 8)])
    snapshot.period_start = date(2026, 1, 7)
    snapshot.assignments = [Assignment(employee_id="e0", demand_id="boundary", fixed=True)]
    if missing == "approval":
        snapshot.employees[0].approvals[0].valid_from = snapshot.period_start
    elif missing == "team":
        snapshot.shifts[0].team_id = "unresolved-source-team"
    else:
        snapshot.shifts[0].kind = "unconfirmed"
    from sp5generator.domain import input_diagnostics
    assert not input_diagnostics(snapshot)
    checked = validate(snapshot, plan(snapshot))
    assert not checked.valid
    assert any(d.code == missing and d.demand_id == "boundary" for d in checked.diagnostics)
    result = solver.solve(snapshot, 3, partial=partial)
    assert result.solver_status == "INFEASIBLE"
    assert not result.assignments
    assert any(d.code == "fixed_conflict" and d.demand_id == "boundary"
               and missing in d.message for d in result.validation.diagnostics)


@pytest.mark.parametrize("partial", [False, True])
def test_weekly_violation_does_not_hide_later_daily_or_weekly_diagnostics(partial):
    snapshot = case(1, [shift("first", 6, 8, 10), shift("second", 13, 8, 10)])
    snapshot.period_end = date(2026, 1, 18)
    snapshot.profiles[0].max_daily_minutes = 9 * 60
    snapshot.profiles[0].max_weekly_minutes = 9 * 60
    checked = validate(snapshot, plan(snapshot))
    assert not checked.valid
    assert [d.date for d in checked.diagnostics if d.code == "daily_limit"] == [
        "2026-01-06", "2026-01-13",
    ]
    assert [d.date for d in checked.diagnostics if d.code == "weekly_limit"] == [
        "2026-01-05", "2026-01-12",
    ]
    result = solver.solve(snapshot, 3, partial=partial)
    assert result.solver_status == ("OPTIMAL" if partial else "INFEASIBLE")
    assert not result.assignments
    if partial:
        assert result.validation.valid
        assert sum(result.vacancies.values()) == 2


@pytest.mark.parametrize("partial", [False, True])
@pytest.mark.parametrize("credit,balance", [(480, 0), (0, 480), (720, -240)])
def test_confirmed_account_adjustments_can_explain_nonselection(partial, credit, balance, tmp_path):
    """A person whose target is covered need not receive a duty.

    These are explicitly configured synthetic values, never imported actual
    account totals or an inference of personal approval.
    """
    from sp5generator.models import Objectives

    snapshot = case(2)
    snapshot.objectives = Objectives(
        hours=1, changes=0, nights=0, weekends=0, holidays=0, wishes=0,
        workday_transitions=0,
    )
    for employee in snapshot.employees:
        employee.target_minutes = 480
    snapshot.employees[0].credit_minutes = credit
    snapshot.employees[0].balance_minutes = balance
    result = solver.solve(snapshot, 3, partial=partial)
    assert result.solver_status == "OPTIMAL" and result.validation.complete
    assert [a.employee_id for a in result.assignments] == ["e1"]
    assert result.metrics["employees"]["e0"]["paid_minutes"] == 0
    assert result.metrics["employees"]["e0"]["deviation_minutes"] == 0
    assert result.metrics["objective_contributions"]["hours"] == 0
    diagnostic = result.metrics["planning_diagnostics"]["employees"]["e0"]
    assert diagnostic["eligible_demands"] == 1
    assert diagnostic["reason"] == "not_selected_with_candidates"
    assert validate(snapshot, result.assignments).complete

    # Exercise the public export gate with the solved plan, not a hand-built
    # success result. Cached metrics must not become an accounting source.
    import csv
    from openpyxl import load_workbook
    from sp5generator.export import export_table

    result.metrics["employees"]["e0"]["deviation_minutes"] = 999999
    csv_path = tmp_path / "account.csv"
    export_table(snapshot, result, csv_path)
    with csv_path.open(encoding="utf-8-sig", newline="") as stream:
        balances = {row[0]: row[1:] for row in csv.reader(stream)
                    if len(row) == 4 and row[0] in {"e0", "e1"}}
    assert balances == {"e0": ["480", str(credit), "0"], "e1": ["480", "480", "0"]}

    xlsx_path = tmp_path / "account.xlsx"
    export_table(snapshot, result, xlsx_path)
    workbook = load_workbook(xlsx_path)
    try:
        sheet = workbook["Stundenübersicht"]
        assert [sheet.cell(5, c).value for c in range(2, 7)] == [
            8, 0, credit / 60, balance / 60, 0,
        ]
        assert [sheet.cell(6, c).value for c in range(2, 7)] == [8, 8, 0, 0, 0]
    finally:
        workbook.close()


@pytest.mark.parametrize("partial", [False, True])
@pytest.mark.parametrize("limit,code", [
    ("max_daily_minutes", "daily_limit"), ("max_weekly_minutes", "weekly_limit"),
])
@pytest.mark.parametrize("credit,balance", [(0, -6000), (6000, 0), (6000, -6000)])
def test_account_adjustments_never_offset_elapsed_hard_limits(partial, limit, code, credit, balance):
    snapshot = case(1)
    employee = snapshot.employees[0]
    employee.target_minutes = 6000
    employee.credit_minutes = credit
    employee.balance_minutes = balance
    setattr(snapshot.profiles[0], limit, 479)
    snapshot.shifts[0].paid_minutes = 60  # Eight hours present, only one paid.
    checked = validate(snapshot, plan(snapshot))
    assert not checked.valid
    assert code in {d.code for d in checked.diagnostics}
    result = solver.solve(snapshot, 3, partial=partial)
    assert result.solver_status == ("OPTIMAL" if partial else "INFEASIBLE")
    assert not result.assignments
    if partial:
        assert result.validation.valid
        assert result.vacancies == {"s": 1}
        assert result.metrics["employees"]["e0"]["deviation_minutes"] == credit + balance - 6000


@pytest.mark.parametrize("partial", [False, True])
@pytest.mark.parametrize("quality_finishes", [False, True])
def test_objective_and_bound_remain_in_their_search_phase(monkeypatch, partial, quality_finishes):
    """No quality gap may be inferred from a feasibility/vacancy fallback."""
    from sp5generator.models import Objectives, Result

    snapshot = case(1)
    snapshot.employees[0].target_minutes = 0
    snapshot.objectives = Objectives(
        hours=2, changes=0, nights=0, weekends=0, holidays=0, wishes=0,
        workday_transitions=0,
    )
    original = cp_model.CpSolver.solve
    calls = []

    def controlled_quality(self, model, *args, **kwargs):
        calls.append(1)
        if len(calls) == 2 and not quality_finishes:
            return cp_model.UNKNOWN
        return original(self, model, *args, **kwargs)

    monkeypatch.setattr(cp_model.CpSolver, "solve", controlled_quality)
    result = solver.solve(snapshot, 3, partial=partial)
    # Exercise the persisted/public JSON contract as well as the live result.
    result = Result.model_validate_json(result.model_dump_json())
    assert len(calls) == 2
    assert result.validation.complete and validate(snapshot, result.assignments).complete
    assert result.metrics["weighted_objective_contributions"]["hours"] == 960
    if quality_finishes:
        assert result.solver_status == "OPTIMAL"
        assert result.metrics["objective_phase"] == "quality"
        assert result.objective_value == result.best_bound == 960
    else:
        assert result.solver_status == "FEASIBLE"
        assert result.parameters["last_optimization_status"] == "UNKNOWN"
        assert result.metrics["objective_phase"] == ("vacancies" if partial else "feasibility")
        if partial:
            assert result.objective_value == result.best_bound == 0
        else:
            assert result.objective_value is None and result.best_bound is None


@pytest.mark.parametrize("partial", [False, True])
@pytest.mark.parametrize("constraint,code", [
    ("weekly", "weekly_limit"), ("rest", "rest"), ("overlap", "overlap"),
])
def test_individual_candidates_do_not_claim_joint_feasibility(partial, constraint, code):
    """An eligible person can be blocked by immutable personal work context."""
    from sp5generator.domain import eligibility
    from sp5generator.models import BoundaryWork

    snapshot = case(2, [shift("new", 7, 8, 8)])
    snapshot.period_start = date(2026, 1, 7)
    profile = snapshot.profiles[0]
    profile.min_rest_minutes = 660
    profile.weekly_rest_minutes = 2160
    if constraint == "weekly":
        previous = shift("context", 5, 8, 8)
        profile.max_weekly_minutes = 480
    elif constraint == "rest":
        previous = shift("context", 6, 22, 8)
    else:
        previous = shift("context", 6, 22, 12)
    snapshot.boundary_work = [BoundaryWork(
        id="context", employee_id="e0", segments=previous.segments, kind="day",
    )]
    assert eligibility(snapshot, snapshot.employees[0], snapshot.demands[0]) == []
    counterfactual = validate(snapshot, [Assignment(employee_id="e0", demand_id="new")])
    assert not counterfactual.valid
    assert code in {d.code for d in counterfactual.diagnostics}

    result = solver.solve(snapshot, 3, partial=partial)
    assert result.solver_status == "OPTIMAL" and result.validation.complete
    assert [(a.employee_id, a.demand_id) for a in result.assignments] == [("e1", "new")]
    assert validate(snapshot, result.assignments).valid
    diagnostics = result.metrics["planning_diagnostics"]
    assert "not a joint feasibility" in diagnostics["semantics"]
    assert diagnostics["employees"]["e0"] == {
        "positive_capacity_demands": 1, "eligible_demands": 1,
        "eligible_required_demands": 1, "exclusions": {},
        "assigned_demands": 0, "reason": "not_selected_with_candidates",
    }


@pytest.mark.parametrize("partial", [False, True])
@pytest.mark.parametrize("planning_day,same_week", [
    (date(2027, 1, 1), True),
    (date(2027, 1, 4), False),
])
@pytest.mark.parametrize("below_limit", [False, True])
def test_iso_year_local_midnight_with_fixed_context(partial, planning_day, same_week, below_limit):
    # January 1 belongs to 2026-W53; January 4 starts 2027-W01.
    # The new duty starts on the preceding UTC date, but consumes exactly
    # 120 minutes on its LOCAL date/week. Fixed context contributes another
    # 120 only in the first case. These expected totals are hand-calculated.
    boundary_day = planning_day - timedelta(days=1)
    snapshot = dated_case(planning_day, planning_day, [
        (localize(boundary_day, "08:00", "Europe/Vienna"),
         localize(boundary_day, "10:00", "Europe/Vienna")),
        (localize(planning_day, "00:30", "Europe/Vienna"),
         localize(planning_day, "02:30", "Europe/Vienna")),
    ])
    profile = snapshot.profiles[0]
    profile.min_rest_minutes = 660
    profile.max_weekly_minutes = (240 if same_week else 120) - int(below_limit)
    snapshot.assignments = [Assignment(employee_id="e0", demand_id="0", fixed=True)]
    snapshot.employees[0].target_minutes = 6000
    for duty in snapshot.shifts:
        duty.paid_minutes = 1
    checked = validate(snapshot, plan(snapshot))
    assert checked.valid is (not below_limit)
    violations = [d for d in checked.diagnostics if d.code == "weekly_limit"]
    assert len(violations) == int(below_limit)
    if violations:
        assert violations[0].date == ("2026-12-28" if same_week else "2027-01-04")
    result = solver.solve(snapshot, 3, partial=partial)
    if below_limit and not partial:
        assert result.solver_status == "INFEASIBLE"
    else:
        assert result.solver_status == "OPTIMAL"
        assert validate(snapshot, result.assignments).valid
        assert {a.demand_id for a in result.assignments} == ({"0"} if below_limit else {"0", "1"})
        assert result.vacancies == ({"1": 1} if below_limit else {})
