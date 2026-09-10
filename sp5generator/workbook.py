"""Readable monthly planning workbooks, built from independently checked results."""

from collections import defaultdict
from datetime import timedelta
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .export import rows, safe_cell, vacancy_counts
from .timeutils import bounds, local_day


INK = "172D35"
TEAL = "117D75"
MUTED = "647680"
LINE = "DEE7E8"
DAY = "E6F3EF"
NIGHT = "EAEFFB"
WEEKEND = "F1F4F6"


def _text(sheet, row, column, value):
    cell = sheet.cell(row, column, safe_cell(value))
    cell.font = Font(name="Calibri", size=10, color=INK)
    cell.alignment = Alignment(vertical="center", wrap_text=True)
    return cell


def _heading(sheet, row, labels):
    for column, label in enumerate(labels, 1):
        cell = _text(sheet, row, column, label)
        cell.font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=INK)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[row].height = 30


def _setup(sheet, title, subtitle, columns):
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.print_options.horizontalCentered = True
    sheet.page_margins.left = sheet.page_margins.right = 0.25
    sheet.oddFooter.center.text = "Seite &P von &N"
    sheet.oddFooter.right.text = "OpenSchichtplaner5 Generator"
    for row, value in ((1, title), (2, subtitle)):
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=max(2, columns))
        cell = _text(sheet, row, 1, value)
        cell.font = Font(name="Calibri", size=18 if row == 1 else 10,
                         bold=row == 1, color=INK if row == 1 else MUTED)
        sheet.row_dimensions[row].height = 32 if row == 1 else 30


def planning_workbook(snapshot, result):
    """Calendar, numeric balances and complete machine-readable assignment detail.

    Calendar cells include every interval on each touched local date. Paid time
    belongs to the shift's starting date, so midnight crossings are counted once.
    All supplied strings remain protected from spreadsheet formula injection.
    """
    workbook = Workbook()
    workbook.remove(workbook.active)
    people = {e.id: e for e in snapshot.employees}
    demands = {d.id: d for d in snapshot.demands}
    shifts = {s.id: s for s in snapshot.shifts}
    positions = {p.id: p for p in snapshot.positions}
    zone = ZoneInfo(snapshot.timezone)
    calendar = defaultdict(list)
    paid = defaultdict(int)
    nights = defaultdict(set)
    workdays = defaultdict(set)
    for assignment in result.assignments:
        demand = demands[assignment.demand_id]
        shift, position = shifts[demand.shift_id], positions[demand.position_id]
        first_day = local_day(bounds(shift)[0], snapshot.timezone)
        if snapshot.period_start <= first_day <= snapshot.period_end:
            paid[assignment.employee_id] += shift.paid_minutes
            if shift.kind == "night":
                nights[assignment.employee_id].add(first_day)
        for interval in shift.segments:
            begin, end = interval.start.astimezone(zone), interval.end.astimezone(zone)
            last = (end.astimezone(ZoneInfo("UTC")) - timedelta(microseconds=1)).astimezone(zone).date()
            day = max(begin.date(), snapshot.period_start)
            while day <= min(last, snapshot.period_end):
                workdays[assignment.employee_id].add(day)
                time_label = f"{begin:%H:%M}–{end:%H:%M}"
                if begin.date() != end.date():
                    time_label += " (+1)" if (end.date() - begin.date()).days == 1 else f" (+{(end.date() - begin.date()).days})"
                label = f"{'◆ ' if assignment.fixed else ''}{shift.name} · {position.name}\n{time_label}"
                calendar[assignment.employee_id, day].append((label, shift.kind))
                day += timedelta(days=1)

    project_name = str(snapshot.metadata.get("project_name") or "Dienstplan")
    status = "Vollständig geprüft" if result.validation.complete else "Teilplan · offener Bedarf"
    subtitle = (f"{snapshot.period_start:%d.%m.%Y} – {snapshot.period_end:%d.%m.%Y}  |  "
                f"{status}  |  {snapshot.timezone}  |  Stand {snapshot.revision}")
    month = snapshot.period_start.replace(day=1)
    weekdays = ("Mo", "Di", "Mi", "Do", "Fr", "Sa", "So")
    while month <= snapshot.period_end:
        following = (month.replace(day=28) + timedelta(days=4)).replace(day=1)
        days = []
        day = max(month, snapshot.period_start)
        while day < following and day <= snapshot.period_end:
            days.append(day)
            day += timedelta(days=1)
        sheet = workbook.create_sheet(f"Plan {month:%Y-%m}")
        _setup(sheet, f"{project_name} · {month:%m/%Y}", subtitle, len(days) + 1)
        _heading(sheet, 4, ["Person", *[f"{weekdays[d.weekday()]}\n{d:%d.%m.}" for d in days]])
        sheet.column_dimensions["A"].width = 27
        for column in range(2, len(days) + 2):
            sheet.column_dimensions[get_column_letter(column)].width = 16
        for row, employee in enumerate(snapshot.employees, 5):
            name = _text(sheet, row, 1, employee.name)
            name.font = Font(name="Calibri", size=10, bold=True, color=INK)
            name.fill = PatternFill("solid", fgColor="F7F9FA")
            max_lines = 1
            for column, day in enumerate(days, 2):
                entries = calendar[employee.id, day]
                label = "\n\n".join(label for label, _ in entries)
                cell = _text(sheet, row, column, label)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                fill = NIGHT if entries and all(kind == "night" for _, kind in entries) else DAY if entries else WEEKEND if day.weekday() >= 5 else "FFFFFF"
                cell.fill = PatternFill("solid", fgColor=fill)
                cell.border = Border(bottom=Side(style="hair", color=LINE), right=Side(style="hair", color=LINE))
                max_lines = max(max_lines, 3 * len(entries))
            sheet.row_dimensions[row].height = min(240, max(50, max_lines * 15))
        legend_row = len(snapshot.employees) + 6
        sheet.merge_cells(start_row=legend_row, start_column=1, end_row=legend_row, end_column=max(2, len(days) + 1))
        _text(sheet, legend_row, 1, "Grün: Tagdienst · Blau: Nachtdienst · ◆: fixiert · (+1): Ende am Folgetag. Übernacht-Dienste erscheinen an allen betroffenen Tagen.")
        sheet.row_dimensions[legend_row].height = 28
        sheet.freeze_panes = "B5"
        sheet.print_title_rows = "1:4"
        sheet.print_title_cols = "A:A"
        sheet.print_area = f"A1:{get_column_letter(len(days) + 1)}{legend_row}"
        month = following

    balances = workbook.create_sheet("Stundenübersicht")
    _setup(balances, "Stundenübersicht", subtitle, 9)
    _heading(balances, 4, ["Person", "Soll (h)", "Geplant (h)", "Gutschrift (h)", "Vortrag (h)", "Saldo (h)", "Arbeitstage", "Nächte", "Beschäftigung (%)"])
    for row, employee in enumerate(snapshot.employees, 5):
        values = [employee.name, employee.target_minutes / 60, paid[employee.id] / 60,
                  employee.credit_minutes / 60, employee.balance_minutes / 60,
                  (paid[employee.id] + employee.credit_minutes + employee.balance_minutes - employee.target_minutes) / 60,
                  len(workdays[employee.id]), len(nights[employee.id]), employee.employment_fraction]
        for column, value in enumerate(values, 1):
            cell = _text(balances, row, column, value)
            if 2 <= column <= 6:
                cell.number_format = '0.00;[Red]-0.00;"–"'
            if row % 2:
                cell.fill = PatternFill("solid", fgColor="F2F7F6")
        balances.row_dimensions[row].height = 24
    balances.freeze_panes = "B5"
    balances.auto_filter.ref = f"A4:I{max(4, len(people) + 4)}"
    balances.column_dimensions["A"].width = 28
    for column in "BCDEFGHI":
        balances.column_dimensions[column].width = 19
    balances.print_title_rows = "1:4"

    gaps = workbook.create_sheet("Offene Stellen")
    _setup(gaps, "Offene Stellen", subtitle, 6)
    _heading(gaps, 4, ["Datum", "Dienst", "Funktion", "Fehlend", "Minimum", "Bedarf-ID"])
    row = 5
    for demand_id, count in vacancy_counts(snapshot, result.assignments).items():
        if not count:
            continue
        demand = demands[demand_id]
        shift, position = shifts[demand.shift_id], positions[demand.position_id]
        values = [str(local_day(bounds(shift)[0], snapshot.timezone)), shift.name, position.name,
                  count, demand.minimum, demand.id]
        for column, value in enumerate(values, 1):
            _text(gaps, row, column, value)
        row += 1
    if row == 5:
        _text(gaps, row, 1, "Kein offener Bedarf.")
    gaps.freeze_panes = "A5"
    for column in "ABCDEF":
        gaps.column_dimensions[column].width = 26

    detail = workbook.create_sheet("Einteilungen")
    detail.sheet_view.showGridLines = False
    headers = {"Person-ID", "Offener Bedarf", "Prüfhinweis"}
    for values in rows(snapshot, result):
        detail.append([safe_cell(value) for value in values])
        if values and values[0] in headers:
            _heading(detail, detail.max_row, values)
    detail.freeze_panes = "C4"
    detail.column_dimensions["B"].width = 30
    for column in "ACDEFGHIJ":
        detail.column_dimensions[column].width = 25
    for row in detail:
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    workbook.active = 0
    workbook.calculation.fullCalcOnLoad = True
    return workbook
